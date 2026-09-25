"""Property-scoped, review-first hotel knowledge and document ingestion.

The SQLite schema is additive. The original operations knowledge tables remain
available for existing manually approved entries and FAQs.
"""

import hashlib
import io
import json
import logging
import re
import sqlite3
import time
import uuid
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from .config import settings

logger = logging.getLogger(__name__)


CATEGORIES = (
    "Hotel Overview", "Rooms", "Dining", "Amenities", "Facilities", "Policies",
    "Transportation", "Parking", "Spa", "Pool", "Gym", "Events", "Promotions",
    "FAQs", "Local Guide", "Attractions", "Emergency", "Contacts", "Guest Services",
    "Accessibility", "Housekeeping", "Check-in / Checkout", "Other",
)
FORMATS = {
    ".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv", ".txt": "text/plain", ".md": "text/markdown",
    ".json": "application/json", ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
}
TEXT_TYPES = {".csv", ".txt", ".md", ".json"}
ZIP_TYPES = {".docx", ".xlsx", ".pptx"}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp"}
STOP = {"the", "and", "for", "with", "what", "when", "where", "have", "this", "that", "hotel", "from", "does", "your", "about"}


def _risk_flags(text: str) -> list[str]:
    lower = text.casefold()
    flags = []
    if any(term in lower for term in ("admin password", "api key", "secret key", "private key", "access token", "client secret")):
        flags.append("possible_secret")
    if any(term in lower for term in ("ignore your system prompt", "ignore previous instructions", "reveal hidden prompt", "expose all admin")):
        flags.append("prompt_injection")
    if any(term in lower for term in ("network diagram", "security procedure", "incident report", "employee record", "pms troubleshooting")):
        flags.append("internal_operations")
    return flags


def safe_filename(name: str) -> str:
    name = name.replace("\\", "/").split("/")[-1]
    name = re.sub(r"[^\w. -]", "_", name, flags=re.UNICODE).strip(" .")[:180]
    if not name or name.startswith("."):
        raise ValueError("Choose a valid filename.")
    return name


def validate_file(name: str, mime: str, data: bytes) -> tuple[str, str]:
    name = safe_filename(name)
    ext = Path(name).suffix.lower()
    if ext not in FORMATS:
        raise ValueError("Unsupported file type. Supported formats: " + ", ".join(FORMATS))
    if len(data) > settings.knowledge_max_file_bytes:
        raise ValueError(f"File exceeds the {settings.knowledge_max_file_bytes // (1024 * 1024)} MB upload limit.")
    if not data:
        raise ValueError("The file is empty.")
    allowed_mimes = {FORMATS[ext], "application/octet-stream"}
    if ext == ".md": allowed_mimes.add("text/plain")
    if ext == ".csv": allowed_mimes.add("application/vnd.ms-excel")
    if mime not in allowed_mimes:
        raise ValueError("File MIME type does not match its extension.")
    magic = {
        ".pdf": data.startswith(b"%PDF-"),
        ".png": data.startswith(b"\x89PNG\r\n\x1a\n"),
        ".jpg": data.startswith(b"\xff\xd8\xff"),
        ".jpeg": data.startswith(b"\xff\xd8\xff"),
        ".webp": data.startswith(b"RIFF") and data[8:12] == b"WEBP",
    }
    if ext in magic and not magic[ext]:
        raise ValueError("File content does not match its extension.")
    if ext in ZIP_TYPES:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 2000 or sum(info.file_size for info in infos) > settings.knowledge_max_extracted_chars * 12:
                    raise ValueError("Document archive exceeds safe extraction limits.")
                if any(info.file_size > max(100, info.compress_size) * 100 for info in infos):
                    raise ValueError("Document archive has an unsafe compression ratio.")
                required = {".docx": "word/document.xml", ".xlsx": "xl/workbook.xml", ".pptx": "ppt/presentation.xml"}[ext]
                if required not in archive.namelist() or any("vbaProject.bin" in info.filename or "/embeddings/" in info.filename or "/activeX/" in info.filename for info in infos):
                    raise ValueError("Invalid Office document or unsupported embedded active content.")
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid Office document archive.") from exc
    if ext in TEXT_TYPES:
        try:
            decoded = data.decode("utf-8-sig")
            if ext == ".json": json.loads(decoded)
            if ext == ".csv":
                import csv
                if any(cell.lstrip().startswith(("=", "+", "@")) for row in csv.reader(io.StringIO(decoded)) for cell in row):
                    raise ValueError("CSV formula cells are not accepted. Export values only.")
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError("Text files must be valid UTF-8; JSON must be valid JSON.") from exc
    return name, ext


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def parse_document(data: bytes, ext: str) -> list[dict[str, Any]]:
    """Return bounded text blocks with stable source locations; never follow links."""
    blocks: list[dict[str, Any]] = []
    if ext in TEXT_TYPES:
        text = data.decode("utf-8-sig")
        if ext == ".json": text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
        if ext == ".csv":
            import csv
            rows = csv.reader(io.StringIO(text))
            headers = next(rows, [])
            if headers: blocks.append({"text": " | ".join(headers), "location": {"row": 1}})
            for row_number, row in enumerate(rows, 2):
                cells = [f"{headers[index]}: {value}" if index < len(headers) and headers[index] else value for index, value in enumerate(row) if value]
                blocks.append({"text": " | ".join(cells), "location": {"row": row_number}})
        else:
            for row_number, line in enumerate(text.splitlines(), 1):
                if line.strip(): blocks.append({"text": line.strip(), "location": {"line": row_number}})
    elif ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data), strict=True)
        if len(reader.pages) > 300: raise ValueError("PDF exceeds the 300 page processing limit.")
        for number, page in enumerate(reader.pages, 1):
            for line in (page.extract_text() or "").splitlines():
                if line.strip(): blocks.append({"text": line.strip(), "location": {"page": number}})
        if not blocks: raise ValueError("PDF contains no machine-readable text. OCR for scanned PDFs is not configured.")
    elif ext == ".docx":
        from docx import Document
        document = Document(io.BytesIO(data))
        for number, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip(): blocks.append({"text": paragraph.text, "location": {"paragraph": number, "section": paragraph.style.name}})
        for table_number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                blocks.append({"text": " | ".join(cell.text for cell in row.cells), "location": {"table": table_number, "row": row_number}})
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book:
                headers: list[str] = []
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [str(value) if value is not None else "" for value in row]
                    if row_number == 1: headers = cells
                    if any(cells):
                        labelled = [f"{headers[index]}: {value}" if row_number > 1 and index < len(headers) and headers[index] else value for index, value in enumerate(cells) if value]
                        blocks.append({"text": " | ".join(labelled), "location": {"sheet": sheet.title, "row": row_number}})
        finally: book.close()
    elif ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        for number, slide in enumerate(presentation.slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    blocks.append({"text": shape.text, "location": {"slide": number}})
                if shape.has_table:
                    for row_number, row in enumerate(shape.table.rows, 1):
                        blocks.append({"text": " | ".join(cell.text for cell in row.cells), "location": {"slide": number, "row": row_number}})
    elif ext in IMAGE_TYPES:
        from PIL import Image, ImageOps
        import pytesseract
        image = Image.open(io.BytesIO(data))
        if image.width * image.height > 30_000_000: raise ValueError("Image exceeds the OCR pixel limit.")
        image = ImageOps.exif_transpose(image)
        try: text = pytesseract.image_to_string(image, timeout=30)
        except (RuntimeError, pytesseract.TesseractNotFoundError) as exc: raise ValueError("OCR processing failed or Tesseract is not installed.") from exc
        for line in text.splitlines():
            if line.strip(): blocks.append({"text": line.strip(), "location": {"image": 1}})
    total = sum(len(block["text"]) for block in blocks)
    if total > settings.knowledge_max_extracted_chars: raise ValueError("Extracted content exceeds the configured processing limit.")
    if not blocks: raise ValueError("Document contains no usable knowledge.")
    return blocks


def _category(text: str) -> str:
    lower = text.casefold()
    terms = {
        "Check-in / Checkout": ("check-in", "check in", "checkout", "check-out", "late checkout"),
        "Dining": ("restaurant", "breakfast", "lunch", "dinner", "menu", "dining"),
        "Pool": ("pool", "swimming"), "Gym": ("gym", "fitness"), "Spa": ("spa", "massage"),
        "Rooms": ("room", "suite", "bed"), "Policies": ("policy", "cancellation", "deposit", "smoking", "pet"),
        "Transportation": ("shuttle", "airport transfer", "transport"), "Parking": ("parking", "valet"),
        "Emergency": ("emergency", "clinic", "hospital"), "Contacts": ("contact", "phone", "email"),
        "Promotions": ("promotion", "offer", "package"), "Events": ("event", "concert"),
        "Accessibility": ("accessible", "wheelchair"), "Housekeeping": ("housekeeping", "laundry"),
    }
    for category, words in terms.items():
        if any(word in lower for word in words): return category
    return "Other"


def _facts(blocks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Conservative extraction: only observed text, never inferred hotel facts."""
    result = []
    section = ""
    for block in blocks:
        raw = block["text"]
        text = _clean(raw)
        if not text: continue
        if (raw.startswith("#") or (len(text) < 80 and text.endswith(":"))):
            section = text.strip("# :")[:120]
            continue
        category = _category(f"{section} {text}")
        title = section or text.split(":", 1)[0][:90]
        title = title or "Document information"
        location = {**block["location"], "section": section} if section else block["location"]
        key, separator, value = text.partition(":")
        structured = {"field": key.strip(), "value": value.strip()} if separator and len(key) <= 80 else {}
        result.append({"title": title[:200], "category": category, "text": text[:4000], "location": location, "structured": structured})
    return result[:1000]


def _tokens(text: str) -> Counter[str]:
    return Counter(word for word in re.findall(r"[\w]+", text.casefold()) if len(word) > 2 and word not in STOP)


class KnowledgeStore:
    def __init__(self, path: Path, upload_root: Path):
        self.path = path
        self.upload_root = upload_root / "knowledge"
        self.upload_root.mkdir(parents=True, exist_ok=True)
        with self._db() as db:
            db.executescript("""
            CREATE TABLE IF NOT EXISTS km_sources (
              source_id TEXT PRIMARY KEY, property_id TEXT NOT NULL, filename TEXT NOT NULL,
              content_type TEXT NOT NULL, extension TEXT NOT NULL, file_bytes INTEGER NOT NULL,
              sha256 TEXT NOT NULL, version INTEGER NOT NULL DEFAULT 1, previous_source_id TEXT,
              status TEXT NOT NULL, error TEXT NOT NULL DEFAULT '', created_by TEXT NOT NULL,
              created_at INTEGER NOT NULL, updated_at INTEGER NOT NULL, processed_at INTEGER,
              FOREIGN KEY(previous_source_id) REFERENCES km_sources(source_id));
            CREATE INDEX IF NOT EXISTS idx_km_sources_property ON km_sources(property_id,status,created_at);
            DROP INDEX IF EXISTS idx_km_sources_digest;
            CREATE UNIQUE INDEX IF NOT EXISTS idx_km_sources_active_digest ON km_sources(property_id,sha256) WHERE status!='archived';
            CREATE TABLE IF NOT EXISTS km_items (
              item_id TEXT PRIMARY KEY, source_id TEXT, property_id TEXT NOT NULL,
              category TEXT NOT NULL, title TEXT NOT NULL, content TEXT NOT NULL,
              structured_json TEXT NOT NULL DEFAULT '{}',
              risk_flags_json TEXT NOT NULL DEFAULT '[]',
              location_json TEXT NOT NULL DEFAULT '{}', visibility TEXT NOT NULL DEFAULT 'admin',
              role_slug TEXT NOT NULL DEFAULT '', status TEXT NOT NULL DEFAULT 'ready_review',
              effective_at INTEGER, expires_at INTEGER, conflict INTEGER NOT NULL DEFAULT 0,
              created_by TEXT NOT NULL, modified_by TEXT NOT NULL, published_by TEXT,
              created_at INTEGER NOT NULL, updated_at INTEGER NOT NULL, published_at INTEGER,
              FOREIGN KEY(source_id) REFERENCES km_sources(source_id) ON DELETE CASCADE);
            CREATE INDEX IF NOT EXISTS idx_km_items_filter ON km_items(property_id,status,visibility,effective_at,expires_at);
            CREATE TABLE IF NOT EXISTS km_chunks (
              chunk_id TEXT PRIMARY KEY, item_id TEXT NOT NULL, property_id TEXT NOT NULL,
              text TEXT NOT NULL, tokens_json TEXT NOT NULL,
              FOREIGN KEY(item_id) REFERENCES km_items(item_id) ON DELETE CASCADE);
            CREATE INDEX IF NOT EXISTS idx_km_chunks_property ON km_chunks(property_id,item_id);
            CREATE TABLE IF NOT EXISTS km_conflicts (
              conflict_id TEXT PRIMARY KEY, property_id TEXT NOT NULL, item_a TEXT NOT NULL,
              item_b TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'open', note TEXT NOT NULL DEFAULT '',
              resolved_by TEXT, resolved_at INTEGER);
            """)
            columns = {row[1] for row in db.execute("PRAGMA table_info(km_items)")}
            if "structured_json" not in columns:
                db.execute("ALTER TABLE km_items ADD COLUMN structured_json TEXT NOT NULL DEFAULT '{}' ")
            if "risk_flags_json" not in columns:
                db.execute("ALTER TABLE km_items ADD COLUMN risk_flags_json TEXT NOT NULL DEFAULT '[]'")

    def _db(self):
        db = sqlite3.connect(self.path, timeout=30)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        return db

    def _file_path(self, property_id: str, source_id: str) -> Path:
        # Only server-generated identifiers become filesystem components.
        return self.upload_root / hashlib.sha256(property_id.encode()).hexdigest()[:20] / source_id

    def upload(self, property_id: str, filename: str, mime: str, data: bytes, actor: str) -> dict[str, Any]:
        filename, ext = validate_file(filename, mime, data)
        now = int(time.time())
        source_id = uuid.uuid4().hex
        with self._db() as db:
            count, used = db.execute("SELECT COUNT(*),COALESCE(SUM(file_bytes),0) FROM km_sources WHERE property_id=?", (property_id,)).fetchone()
            if count >= settings.knowledge_max_files_per_property: raise ValueError("Property document limit reached.")
            if used + len(data) > settings.knowledge_storage_quota_bytes: raise ValueError("Property storage quota exceeded.")
            duplicate = db.execute("SELECT source_id FROM km_sources WHERE property_id=? AND sha256=? AND status!='archived'", (property_id, hashlib.sha256(data).hexdigest())).fetchone()
            if duplicate: raise ValueError("This document is already uploaded for this property.")
            path = self._file_path(property_id, source_id)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
            try:
                db.execute("INSERT INTO km_sources (source_id,property_id,filename,content_type,extension,file_bytes,sha256,status,created_by,created_at,updated_at) VALUES (?,?,?,?,?,?,?,'uploaded',?,?,?)",
                           (source_id, property_id, filename, mime, ext, len(data), hashlib.sha256(data).hexdigest(), actor, now, now))
            except sqlite3.IntegrityError as exc:
                path.unlink(missing_ok=True)
                raise ValueError("This document is already uploaded for this property.") from exc
            except Exception:
                path.unlink(missing_ok=True)
                raise
        logger.info("knowledge_upload property=%s source=%s bytes=%d", property_id, source_id, len(data))
        return self.get_source(property_id, source_id)

    def get_source(self, property_id: str, source_id: str) -> dict[str, Any] | None:
        with self._db() as db:
            row = db.execute("SELECT * FROM km_sources WHERE property_id=? AND source_id=?", (property_id, source_id)).fetchone()
        return dict(row) if row else None

    def list_sources(self, property_id: str) -> list[dict[str, Any]]:
        with self._db() as db:
            rows = db.execute("SELECT * FROM km_sources WHERE property_id=? ORDER BY created_at DESC", (property_id,)).fetchall()
        return [dict(row) for row in rows]

    def process(self, property_id: str, source_id: str) -> dict[str, Any]:
        source = self.get_source(property_id, source_id)
        if not source: raise KeyError("Document not found.")
        if source["status"] in {"ready_review", "conflict_detected", "published", "archived"}: return source
        now = int(time.time())
        with self._db() as db:
            claim = db.execute("UPDATE km_sources SET status='processing',error='',updated_at=? WHERE property_id=? AND source_id=? AND (status IN ('uploaded','processing_failed') OR (status='processing' AND updated_at<?))", (now, property_id, source_id, now - 300))
            if not claim.rowcount:
                return self.get_source(property_id, source_id)
        try:
            blocks = parse_document(self._file_path(property_id, source_id).read_bytes(), source["extension"])
            facts = _facts(blocks)
            if not facts: raise ValueError("Document contains no usable knowledge.")
            now = int(time.time())
            with self._db() as db:
                db.execute("DELETE FROM km_items WHERE property_id=? AND source_id=?", (property_id, source_id))
                for fact in facts:
                    item_id = uuid.uuid4().hex
                    db.execute("INSERT INTO km_items (item_id,source_id,property_id,category,title,content,structured_json,risk_flags_json,location_json,created_by,modified_by,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                               (item_id, source_id, property_id, fact["category"], fact["title"], fact["text"], json.dumps(fact["structured"]), json.dumps(_risk_flags(fact["text"])), json.dumps(fact["location"]), source["created_by"], source["created_by"], now, now))
                    # One bounded, location-preserving chunk per extracted fact.
                    db.execute("INSERT INTO km_chunks VALUES (?,?,?,?,?)", (uuid.uuid4().hex, item_id, property_id, fact["text"], json.dumps(_tokens(fact["text"]))))
                db.execute("UPDATE km_sources SET status='ready_review',processed_at=?,updated_at=? WHERE property_id=? AND source_id=?", (now, now, property_id, source_id))
            self.detect_conflicts(property_id, source_id)
            logger.info("knowledge_processed property=%s source=%s items=%d", property_id, source_id, len(facts))
        except Exception as exc:
            with self._db() as db:
                db.execute("UPDATE km_sources SET status='processing_failed',error=?,updated_at=? WHERE property_id=? AND source_id=?", (str(exc)[:400], int(time.time()), property_id, source_id))
            logger.warning("knowledge_processing_failed property=%s source=%s error_type=%s", property_id, source_id, exc.__class__.__name__)
        return self.get_source(property_id, source_id)

    def pending_sources(self) -> list[tuple[str, str]]:
        now = int(time.time())
        with self._db() as db:
            rows = db.execute("SELECT property_id,source_id FROM km_sources WHERE status='uploaded' OR (status='processing' AND updated_at<?) LIMIT 20", (now - 300,)).fetchall()
        return [(row["property_id"], row["source_id"]) for row in rows]

    def list_items(self, property_id: str, *, source_id: str | None = None, status: str | None = None, query: str = "") -> list[dict[str, Any]]:
        sql = "SELECT * FROM km_items WHERE property_id=?"
        params: list[Any] = [property_id]
        if source_id: sql += " AND source_id=?"; params.append(source_id)
        if status: sql += " AND status=?"; params.append(status)
        if query: sql += " AND (title LIKE ? OR content LIKE ?)"; params.extend([f"%{query[:100]}%"] * 2)
        sql += " ORDER BY updated_at DESC LIMIT 1000"
        with self._db() as db: rows = db.execute(sql, params).fetchall()
        return [self._item(row) for row in rows]

    @staticmethod
    def _item(row: sqlite3.Row) -> dict[str, Any]:
        result = dict(row)
        result["location"] = json.loads(result.pop("location_json"))
        result["structured"] = json.loads(result.pop("structured_json"))
        result["risk_flags"] = json.loads(result.pop("risk_flags_json"))
        result["conflict"] = bool(result["conflict"])
        return result

    def get_item(self, property_id: str, item_id: str) -> dict[str, Any] | None:
        with self._db() as db: row = db.execute("SELECT * FROM km_items WHERE property_id=? AND item_id=?", (property_id, item_id)).fetchone()
        return self._item(row) if row else None

    def create_draft(self, property_id: str, title: str, content: str, category: str, actor: str) -> dict[str, Any]:
        title, content = title.strip()[:200], content.strip()[:4000]
        if not title or not content: raise ValueError("Title and content are required.")
        if category not in CATEGORIES: raise ValueError("Choose a supported category.")
        now, item_id = int(time.time()), uuid.uuid4().hex
        key, separator, value = content.partition(":")
        structured = {"field": key.strip(), "value": value.strip()} if separator and len(key) <= 80 else {}
        with self._db() as db:
            db.execute("INSERT INTO km_items (item_id,source_id,property_id,category,title,content,structured_json,risk_flags_json,location_json,created_by,modified_by,created_at,updated_at) VALUES (?,NULL,?,?,?,?,?,?,?,?,?,?,?)",
                       (item_id, property_id, category, title, content, json.dumps(structured), json.dumps(_risk_flags(content)), json.dumps({"source": "Admin Knowledge Entry"}), actor, actor, now, now))
            db.execute("INSERT INTO km_chunks VALUES (?,?,?,?,?)", (uuid.uuid4().hex, item_id, property_id, content, json.dumps(_tokens(content))))
        return self.get_item(property_id, item_id)

    def update_item(self, property_id: str, item_id: str, changes: dict[str, Any], actor: str) -> dict[str, Any]:
        current = self.get_item(property_id, item_id)
        if not current: raise KeyError("Knowledge item not found.")
        if current["status"] == "published": raise ValueError("Unpublish this item before editing it.")
        category = str(changes.get("category", current["category"]))[:80]
        if category not in CATEGORIES: raise ValueError("Choose a supported category.")
        visibility = str(changes.get("visibility", current["visibility"]))
        if visibility not in {"guest", "admin", "manager", "engineering", "staff", "role"}: raise ValueError("Invalid visibility.")
        role_slug = str(changes.get("role_slug", current["role_slug"]))[:80] if visibility == "role" else ""
        if visibility == "role" and not role_slug: raise ValueError("Choose a role for custom visibility.")
        content = str(changes.get("content", current["content"])).strip()[:4000]
        title = str(changes.get("title", current["title"])).strip()[:200]
        if not title or not content: raise ValueError("Title and content are required.")
        flags = _risk_flags(content)
        if visibility == "guest" and flags:
            raise ValueError("This item contains potentially sensitive or instructive text. Remove it before guest publication.")
        key, separator, value = content.partition(":")
        structured = {"field": key.strip(), "value": value.strip()} if separator and len(key) <= 80 else {}
        effective = changes.get("effective_at", current["effective_at"])
        expires = changes.get("expires_at", current["expires_at"])
        if effective is not None: effective = int(effective)
        if expires is not None: expires = int(expires)
        if effective and expires and expires <= effective: raise ValueError("Expiration must follow the effective date.")
        with self._db() as db:
            db.execute("UPDATE km_items SET category=?,title=?,content=?,structured_json=?,risk_flags_json=?,visibility=?,role_slug=?,effective_at=?,expires_at=?,modified_by=?,updated_at=? WHERE property_id=? AND item_id=?",
                       (category, title, content, json.dumps(structured), json.dumps(flags), visibility, role_slug, effective, expires, actor, int(time.time()), property_id, item_id))
            db.execute("UPDATE km_chunks SET text=?,tokens_json=? WHERE property_id=? AND item_id=?", (content, json.dumps(_tokens(content)), property_id, item_id))
        return self.get_item(property_id, item_id)

    def set_status(self, property_id: str, item_id: str, status: str, actor: str) -> dict[str, Any]:
        if status not in {"approved", "published", "ready_review", "archived", "rejected"}: raise ValueError("Invalid knowledge status.")
        item = self.get_item(property_id, item_id)
        if not item: raise KeyError("Knowledge item not found.")
        if status == "published" and item["status"] != "approved": raise ValueError("Approve this item before publishing.")
        if status == "published" and item["conflict"]: raise ValueError("Resolve the conflict before publishing.")
        if status == "published" and item["visibility"] == "guest" and item["risk_flags"]:
            raise ValueError("Potentially sensitive text cannot be published to Guest AI.")
        now = int(time.time())
        with self._db() as db:
            db.execute("UPDATE km_items SET status=?,modified_by=?,updated_at=?,published_by=CASE WHEN ?='published' THEN ? ELSE published_by END,published_at=CASE WHEN ?='published' THEN ? ELSE published_at END WHERE property_id=? AND item_id=?",
                       (status, actor, now, status, actor, status, now, property_id, item_id))
        return self.get_item(property_id, item_id)

    def search(self, property_id: str, query: str, *, guest: bool, role_slug: str = "", limit: int = 5) -> list[dict[str, Any]]:
        terms = _tokens(query)
        if not terms: return []
        now = int(time.time())
        sql = """SELECT i.*,s.filename,c.text,c.tokens_json FROM km_items i
          LEFT JOIN km_sources s ON s.source_id=i.source_id AND s.property_id=i.property_id
          JOIN km_chunks c ON c.item_id=i.item_id AND c.property_id=i.property_id
          WHERE i.property_id=?"""
        params: list[Any] = [property_id]
        if guest:
            sql += " AND i.status='published' AND i.conflict=0 AND (i.effective_at IS NULL OR i.effective_at<=?) AND (i.expires_at IS NULL OR i.expires_at>?)"
            params.extend([now, now])
        else:
            sql += " AND i.status IN ('published','approved','ready_review')"
        if guest: sql += " AND i.visibility='guest'"
        elif role_slug != "super-admin":
            sql += " AND (i.visibility IN ('guest','staff') OR (i.visibility='admin' AND ? IN ('property-administrator','content-manager')) OR (i.visibility='manager' AND ? IN ('property-administrator','property-manager')) OR (i.visibility='engineering' AND ?='engineering') OR (i.visibility='role' AND i.role_slug=?))"
            params.extend([role_slug, role_slug, role_slug, role_slug])
        with self._db() as db: rows = db.execute(sql, params).fetchall()
        scored = []
        for row in rows:
            tokens = json.loads(row["tokens_json"])
            score = sum(min(tokens.get(term, 0), 3) for term in terms)
            score += sum(2 for term in terms if term in row["title"].casefold())
            if score:
                scored.append((score, {"title": row["title"], "answer": row["text"], "category": row["category"],
                                       "source": row["filename"] or "Admin Knowledge Entry", "source_id": row["source_id"],
                                       "item_id": row["item_id"], "status": row["status"], "location": json.loads(row["location_json"])}))
        scored.sort(key=lambda pair: pair[0], reverse=True)
        logger.info("knowledge_search property=%s guest=%s matches=%d", property_id, guest, len(scored))
        return [item for _, item in scored[:limit]]

    def detect_conflicts(self, property_id: str, source_id: str) -> list[dict[str, Any]]:
        # Conservative conflict detection on explicit key:value statements.
        with self._db() as db:
            new_rows = db.execute("SELECT * FROM km_items WHERE property_id=? AND source_id=?", (property_id, source_id)).fetchall()
            old_rows = db.execute("SELECT * FROM km_items WHERE property_id=? AND source_id!=? AND status IN ('published','approved','ready_review')", (property_id, source_id)).fetchall()
            for new in new_rows:
                key, sep, value = new["content"].partition(":")
                if not sep or len(key) > 80: continue
                for old in old_rows:
                    old_key, old_sep, old_value = old["content"].partition(":")
                    if old_sep and key.strip().casefold() == old_key.strip().casefold() and value.strip().casefold() != old_value.strip().casefold():
                        exists = db.execute("SELECT 1 FROM km_conflicts WHERE property_id=? AND item_a=? AND item_b=?", (property_id, old["item_id"], new["item_id"])).fetchone()
                        if not exists:
                            db.execute("INSERT INTO km_conflicts (conflict_id,property_id,item_a,item_b) VALUES (?,?,?,?)", (uuid.uuid4().hex, property_id, old["item_id"], new["item_id"]))
                        db.execute("UPDATE km_items SET conflict=1 WHERE property_id=? AND item_id IN (?,?)", (property_id, old["item_id"], new["item_id"]))
            if db.execute("SELECT 1 FROM km_items WHERE property_id=? AND source_id=? AND conflict=1", (property_id, source_id)).fetchone():
                db.execute("UPDATE km_sources SET status='conflict_detected' WHERE property_id=? AND source_id=?", (property_id, source_id))
        return self.conflicts(property_id)

    def conflicts(self, property_id: str) -> list[dict[str, Any]]:
        with self._db() as db:
            rows = db.execute("SELECT * FROM km_conflicts WHERE property_id=? ORDER BY status, rowid DESC", (property_id,)).fetchall()
        return [dict(row) for row in rows]

    def resolve_conflict(self, property_id: str, conflict_id: str, winner_id: str, note: str, actor: str) -> dict[str, Any]:
        with self._db() as db:
            row = db.execute("SELECT * FROM km_conflicts WHERE property_id=? AND conflict_id=?", (property_id, conflict_id)).fetchone()
            if not row: raise KeyError("Conflict not found.")
            if winner_id not in {row["item_a"], row["item_b"]}: raise ValueError("Choose one of the conflicting items.")
            loser = row["item_b"] if winner_id == row["item_a"] else row["item_a"]
            db.execute("UPDATE km_items SET status='archived',conflict=0 WHERE property_id=? AND item_id=?", (property_id, loser))
            db.execute("UPDATE km_items SET conflict=0 WHERE property_id=? AND item_id=?", (property_id, winner_id))
            db.execute("UPDATE km_conflicts SET status='resolved',note=?,resolved_by=?,resolved_at=? WHERE property_id=? AND conflict_id=?", (note[:1000], actor, int(time.time()), property_id, conflict_id))
            db.execute("UPDATE km_sources SET status='ready_review' WHERE property_id=? AND status='conflict_detected' AND source_id IN (SELECT source_id FROM km_items WHERE item_id=?)", (property_id, winner_id))
        return {"conflict_id": conflict_id, "winner_id": winner_id, "archived_id": loser}

    def replace(self, property_id: str, previous_id: str, filename: str, mime: str, data: bytes, actor: str) -> dict[str, Any]:
        previous = self.get_source(property_id, previous_id)
        if not previous: raise KeyError("Document not found.")
        fresh = self.upload(property_id, filename, mime, data, actor)
        with self._db() as db:
            db.execute("UPDATE km_sources SET previous_source_id=?,version=? WHERE property_id=? AND source_id=?", (previous_id, previous["version"] + 1, property_id, fresh["source_id"]))
        return self.get_source(property_id, fresh["source_id"])

    def supersede(self, property_id: str, new_id: str) -> None:
        new = self.get_source(property_id, new_id)
        if not new or not new["previous_source_id"]: raise ValueError("This document has no previous version.")
        if new["status"] not in {"ready_review", "published"}: raise ValueError("Process the replacement first.")
        old_id = new["previous_source_id"]
        with self._db() as db:
            db.execute("UPDATE km_items SET status='archived' WHERE property_id=? AND source_id=?", (property_id, old_id))
            db.execute("UPDATE km_sources SET status='archived' WHERE property_id=? AND source_id=?", (property_id, old_id))

    def delete_source(self, property_id: str, source_id: str) -> bool:
        source = self.get_source(property_id, source_id)
        if not source: return False
        with self._db() as db:
            db.execute("DELETE FROM km_conflicts WHERE property_id=? AND (item_a IN (SELECT item_id FROM km_items WHERE property_id=? AND source_id=?) OR item_b IN (SELECT item_id FROM km_items WHERE property_id=? AND source_id=?))", (property_id, property_id, source_id, property_id, source_id))
            db.execute("UPDATE km_sources SET previous_source_id=NULL WHERE property_id=? AND previous_source_id=?", (property_id, source_id))
            db.execute("DELETE FROM km_sources WHERE property_id=? AND source_id=?", (property_id, source_id))
            db.execute("UPDATE km_items SET conflict=0 WHERE property_id=? AND conflict=1 AND item_id NOT IN (SELECT item_a FROM km_conflicts WHERE property_id=? AND status='open' UNION SELECT item_b FROM km_conflicts WHERE property_id=? AND status='open')", (property_id, property_id, property_id))
        self._file_path(property_id, source_id).unlink(missing_ok=True)
        return True

    def health(self, property_id: str) -> dict[str, Any]:
        now = int(time.time())
        required = ("Rooms", "Dining", "Amenities", "Policies", "Transportation", "Emergency", "Contacts", "Check-in / Checkout")
        fields = {
            "check-in time": ("check-in", "check in"),
            "checkout time": ("checkout", "check-out"),
            "breakfast hours": ("breakfast",),
            "airport transfer": ("airport transfer", "airport shuttle"),
            "pet policy": ("pet policy", "pets allowed", "no pets"),
            "emergency contact": ("emergency contact", "emergency phone"),
        }
        with self._db() as db:
            rows = db.execute("SELECT category,content,status,effective_at,expires_at FROM km_items WHERE property_id=?", (property_id,)).fetchall()
            pending = db.execute("SELECT COUNT(*) FROM km_items WHERE property_id=? AND status='ready_review'", (property_id,)).fetchone()[0]
            conflicts = db.execute("SELECT COUNT(*) FROM km_conflicts WHERE property_id=? AND status='open'", (property_id,)).fetchone()[0]
        active = [row for row in rows if row["status"] == "published" and (row["effective_at"] is None or row["effective_at"] <= now) and (row["expires_at"] is None or row["expires_at"] > now)]
        covered = {row["category"] for row in active}
        missing = [category for category in required if category not in covered]
        text = "\n".join(row["content"].casefold() for row in active)
        missing_fields = [field for field, terms in fields.items() if not any(term in text for term in terms)]
        score = round(100 * (len(required) + len(fields) - len(missing) - len(missing_fields)) / (len(required) + len(fields)))
        return {"coverage_percent": score, "covered_categories": sorted(covered & set(required)),
                "required_categories": list(required), "missing_categories": missing, "pending_review": pending,
                "missing_fields": missing_fields,
                "open_conflicts": conflicts, "expired_items": sum(row["expires_at"] is not None and row["expires_at"] <= now for row in rows),
                "methodology": "Equal-weight coverage of eight required categories and six named core facts, counted only from published active items. Keyword presence is a proxy, not verification of completeness."}
